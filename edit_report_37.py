from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path(r"E:\Hoc_Tap\Đồ Án Cơ Sở\Báo cáo đồ án cơ sở.docx")
OUTPUT = Path(r"C:\Users\Administrator\Documents\Đồ án cơ sở\Báo cáo đồ án cơ sở - đã sửa mục 3.7.docx")
IMAGES = Path(r"C:\Users\Administrator\Documents\Đồ án cơ sở\screenshots_web")

SECTIONS = [
    (
        "3.7.1. Giao diện trang chủ và danh sách sản phẩm",
        "Màn hình trang chủ hiển thị danh sách hàng hóa theo dạng lưới, hỗ trợ tìm kiếm, "
        "lọc theo nhóm sản phẩm và truy cập nhanh các chức năng chính của hệ thống.",
        "mobile_02_customer_home.png",
        "Hình 3.22. Giao diện trang chủ và danh sách sản phẩm",
    ),
    (
        "3.7.2. Giao diện menu chức năng",
        "Menu điều hướng cung cấp lối truy cập đến thông tin cá nhân và các chức năng "
        "liên quan đến tài khoản người dùng.",
        "mobile_09_navigation_menu.png",
        "Hình 3.23. Giao diện menu chức năng",
    ),
    (
        "3.7.3. Giao diện chi tiết sản phẩm",
        "Màn hình chi tiết trình bày hình ảnh, mã sản phẩm, danh mục, đơn vị tính, giá bán, "
        "tồn kho, mô tả và khu vực đánh giá của khách hàng.",
        "mobile_10_product_detail.png",
        "Hình 3.24. Giao diện chi tiết sản phẩm",
    ),
    (
        "3.7.4. Giao diện quản lý voucher",
        "Màn hình voucher cho phép khách hàng xem các mã ưu đãi hiện có, lưu voucher vào "
        "tài khoản và sử dụng khi thanh toán.",
        "mobile_06_customer_vouchers.png",
        "Hình 3.25. Giao diện quản lý voucher",
    ),
    (
        "3.7.5. Giao diện giỏ hàng và thanh toán",
        "Màn hình giỏ hàng cho phép kiểm tra sản phẩm, điều chỉnh số lượng, xem tổng tiền, "
        "chọn hình thức nhận hàng và tiến hành đặt hàng.",
        "mobile_07_customer_cart.png",
        "Hình 3.26. Giao diện giỏ hàng và thanh toán",
    ),
    (
        "3.7.6. Giao diện theo dõi đơn hàng",
        "Màn hình đơn hàng hiển thị mã đơn, thời gian đặt, hình thức nhận hàng, trạng thái "
        "thanh toán, danh sách sản phẩm và tổng giá trị của từng đơn.",
        "mobile_08_customer_orders.png",
        "Hình 3.27. Giao diện theo dõi đơn hàng",
    ),
    (
        "3.7.7. Giao diện tài khoản khách hàng",
        "Màn hình tài khoản tổng hợp thông tin cá nhân, điểm thành viên, tổng chi tiêu, "
        "thông tin liên hệ và các lối tắt đến voucher, đơn hàng của khách hàng.",
        "mobile_05_customer_profile.png",
        "Hình 3.28. Giao diện tài khoản khách hàng",
    ),
]


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    elif not value and node is not None:
        p_pr.remove(node)


def insert_before(anchor, paragraph):
    anchor._p.addprevious(paragraph._p)


def add_before(doc, anchor, text="", style=None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    insert_before(anchor, paragraph)
    return paragraph


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for _, _, image, _ in SECTIONS:
        if not (IMAGES / image).exists():
            raise FileNotFoundError(IMAGES / image)

    temp = OUTPUT.with_suffix(".working.docx")
    shutil.copy2(SOURCE, temp)
    doc = Document(temp)

    heading = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("3.7. Thiết kế giao diện người dùng hệ thống")
    )
    chapter4 = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("CHƯƠNG 4")
    )

    current = heading._p.getnext()
    while current is not None and current is not chapter4._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following

    intro = add_before(
        doc,
        chapter4,
        "Phần này trình bày các giao diện tiêu biểu của ứng dụng quản lý và bán lẻ "
        "siêu thị trên thiết bị di động. Các màn hình được nhóm theo chức năng để thể "
        "hiện rõ quá trình người dùng tìm kiếm sản phẩm, mua hàng và quản lý tài khoản.",
    )
    intro.paragraph_format.space_after = Pt(6)
    intro.paragraph_format.line_spacing = 1.15

    for title, description, image_name, caption in SECTIONS:
        subheading = add_before(doc, chapter4, title, style="Heading 3")
        subheading.paragraph_format.space_before = Pt(8)
        subheading.paragraph_format.space_after = Pt(3)
        set_keep_with_next(subheading)

        body = add_before(doc, chapter4, description)
        body.paragraph_format.space_after = Pt(4)
        body.paragraph_format.line_spacing = 1.15
        set_keep_with_next(body)

        picture_p = doc.add_paragraph()
        picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_p.paragraph_format.space_before = Pt(2)
        picture_p.paragraph_format.space_after = Pt(2)
        picture_p.add_run().add_picture(str(IMAGES / image_name), height=Inches(5.6))
        set_keep_with_next(picture_p)
        insert_before(chapter4, picture_p)

        caption_p = add_before(doc, chapter4, caption)
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.space_after = Pt(6)
        for run in caption_p.runs:
            run.italic = True
            run.font.size = Pt(11)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    temp.unlink(missing_ok=True)
    print(OUTPUT)


if __name__ == "__main__":
    main()
